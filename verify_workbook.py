#!/usr/bin/env python3
# verify_workbook.py — Verify recalculation and GATE-VAL logic
from openpyxl import load_workbook

wb = load_workbook("out/recalc/HERMES-REO_Model_Suite.xlsx", data_only=True)

print("=== RECALCULATION VERIFICATION ===\n")

# Fix-and-Flip
ws = wb["Fix-and-Flip"]
print("FIX-AND-FLIP:")
for cell, label in [("B10","Total Uses"),("B13","HM loan"),("B14","Equity"),
                    ("B31","Total interest"),("B42","Net profit"),("B44","ROI")]:
    v = ws[cell].value
    print(f"  {label}: {v}")

# Comp & ARV
ws2 = wb["Comp & ARV"]
print(f"\nCOMP & ARV:")
for cell, label in [("B14","Weighted ARV"),("B15","ARV used"),("B18","Comp source"),
                    ("B19","Source rigor"),("B20","Confidence"),("B21","Active floor"),
                    ("B22","GATE-VAL verdict"),("B23","GATE-VAL pass")]:
    v = ws2[cell].value
    print(f"  {label}: {v}")

# Deal Screener
ws3 = wb["Deal Screener"]
print(f"\nDEAL SCREENER:")
print(f"  MAO: {ws3['B10'].value}")
print(f"  1% rule: {ws3['B12'].value}")
print(f"  Verdict: {ws3['B17'].value}")

# BRRRR
ws4 = wb["BRRRR"]
print(f"\nBRRRR:")
print(f"  All-in: {ws4['B4'].value}")
print(f"  Refi loan: {ws4['B16'].value}")
print(f"  Cash left in: {ws4['B19'].value}")
print(f"  % recycled: {ws4['B20'].value}")
print(f"  Verdict: {ws4['B28'].value}")

# Tax-Deed
ws5 = wb["Tax-Deed"]
print(f"\nTAX-DEED:")
print(f"  Premium yr1: {ws5['B14'].value}")
print(f"  Guaranteed return: {ws5['B17'].value}")
print(f"  Annualized yield: {ws5['B18'].value}")

# Sensitivity
ws6 = wb["Sensitivity"]
print(f"\nSENSITIVITY:")
print(f"  Center (base): {ws6.cell(row=8, column=3).value}")
print(f"  Top-left (-10%/-10%): {ws6.cell(row=6, column=2).value}")
print(f"  Bottom-right (+10%/+20%): {ws6.cell(row=10, column=5).value}")

# Checks
ws7 = wb["Checks"]
print(f"\nCHECKS:")
print(f"  Recalc self-test: {ws7['B6'].value}")
print(f"  S&U ties: {ws7['B13'].value}")

# Named ranges
wb2 = load_workbook("out/HERMES-REO_Model_Suite.xlsx")
names = sorted(wb2.defined_names)
print(f"\nNAMED RANGES: {len(names)}")
for n in names[:20]:
    print(f"  {n}: {wb2.defined_names[n].attr_text}")
print(f"  ... ({len(names)} total)")

# Tab list
print(f"\nTABS ({len(wb2.sheetnames)}): {', '.join(wb2.sheetnames)}")