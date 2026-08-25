"""Structural workbook regression tests (audit C4/C5/C6/H8).

These assert the FORMULAS are correct in the generated workbook. Computed
values are covered by the recalc-marked suite in test_recalc.py.
"""
import os

import pytest
from openpyxl import load_workbook

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(REPO, "out", "HERMES-REO_Model_Suite.xlsx")

pytestmark = pytest.mark.skipif(
    not os.path.exists(XLSX),
    reason="workbook not built yet — run `make build`",
)


def _wb():
    return load_workbook(XLSX)


def test_brrrr_cash_flow_adds_negative_debt_service():
    """AUDIT C4: '=B12-B13' added debt service (PMT is negative)."""
    ws = _wb()["BRRRR"]
    assert ws["B24"].value.endswith("+B13")


def test_weighted_arv_uses_full_filtered_range():
    """AUDIT C5: MEDIAN must span rows 5..10 via the zero-filter helper col."""
    ws = _wb()["Comp & ARV"]
    weighted = None
    for row in ws.iter_rows(min_col=1, max_col=2):
        if row[0].value and "Weighted ARV" in str(row[0].value):
            weighted = row[1].value
    assert weighted, "weighted ARV cell not found"
    assert "MEDIAN(L5:L10)" in weighted


def test_brrrr_verdict_checks_dscr_before_infinite_return():
    """AUDIT C6: DSCR gate must precede the infinite-return branch."""
    formula = _wb()["BRRRR"]["B28"].value
    assert formula.find("min_dscr") < formula.find("B19<=0")


def test_hard_money_loan_is_gated():
    """AUDIT H8: loan sizing must use gated arv_used, not raw assumption."""
    formula = _wb()["Fix-and-Flip"]["B13"].value
    assert "gate_val_pass" in formula and "arv_used" in formula


def test_no_cellis_containsText_rules_remain():
    """AUDIT M11: CellIsRule containsText is ignored by Excel — all replaced."""
    wb = _wb()
    bad = []
    for ws in wb.worksheets:
        for rng in ws.conditional_formatting:
            for rule in rng.rules:
                if rule.type == "cellIs" and rule.operator == "containsText":
                    bad.append(f"{ws.title}!{rng.sqref}")
    assert not bad, f"invalid conditional formats remain: {bad}"


def test_all_15_tabs_present():
    expected = {"README", "Assumptions", "Comp & ARV", "Rehab Estimator",
                "Deal Screener", "Fix-and-Flip", "BRRRR", "Rental DCF",
                "Tax-Deed", "Wholesale", "Auction Bid", "Subject-To",
                "Sensitivity", "Portfolio", "Checks"}
    assert set(_wb().sheetnames) == expected
